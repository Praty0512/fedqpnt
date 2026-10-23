"""Convert the conference paper (main.tex + main.aux/main.bbl from a LaTeX build) to an IEEE-style Word file.

Handles only the constructs used in main.tex. Usage: python tex2docx.py  ->  main.docx
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
FIG = HERE.parents[1] / "results" / "m4r3" / "figures"
tex = (HERE / "main.tex").read_text(encoding="utf-8")
tex = re.sub(r"(?m)^%.*$", "", tex)
aux = (HERE / "main.aux").read_text(encoding="utf-8")
num = {k: n for k, n in re.findall(r"\\bibcite\{([^}]+)\}\{(\d+)\}", aux)}

SYM = {r"\kappa": "κ", r"\sigma": "σ", r"\tau": "τ", r"\psi": "ψ", r"\ln": " ln", r"\lVert": "‖", r"\rVert": "‖",
       r"\mu": "µ", r"\times": "×", r"\to": "→", r"\approx": "≈", r"\tfrac12": "½", r"\le": "≤"}


def math(m):
    s = m
    s = re.sub(r"\\text(?:rm)?\{([^}]*)\}", r"\1", s)
    s = re.sub(r"\\mathrm\{([^}]*)\}", r"\1", s)
    for k, v in SYM.items():
        s = s.replace(k, v)
    s = re.sub(r"_\{([^}]*)\}", r"_\1", s)
    s = re.sub(r"\^\{([^}]*)\}", r"^\1", s)
    s = s.replace("\\!", "").replace("\\,", " ").replace("{", "").replace("}", "")
    s = re.sub(r"10\^(-?\d+)", lambda x: "10" + x.group(1).translate(str.maketrans("-0123456789", "⁻⁰¹²³⁴⁵⁶⁷⁸⁹")), s)
    s = re.sub(r"_([A-Za-z0-9]+)", lambda x: "\x01" + x.group(1) + "\x02", s)   # subscript markers
    return s


ACC = {"'": "́", "`": "̀", '"': "̈", "v": "̌", "^": "̂", "~": "̃", "c": "̧"}


def accents(s):
    import unicodedata
    s = re.sub(r"\\(?:(['`\"^~])\{?\\?([A-Za-z])\}?|([vc])\{\\?([A-Za-z])\})", lambda m: unicodedata.normalize("NFC", (m.group(2) or m.group(4)) + ACC[m.group(1) or m.group(3)]), s)
    return s


def clean(s):
    s = accents(s)
    s = re.sub(r"\\hskip[^\\]*?minus\s*[\d.]+em", "", s)
    s = re.sub(r"1em plus 0\.5em minus 0\.4em", "", s)
    s = re.sub(r"\\cite\{([^}]+)\}", lambda m: ", ".join(f"[{num[k.strip()]}]" for k in m.group(1).split(",")), s)
    s = re.sub(r"\$([^$]+)\$", lambda m: math(m.group(1)), s)
    s = re.sub(r"Table~\\ref\{[^}]+\}", "Table I", s)
    s = re.sub(r"Fig\.~\\ref\{fig:clock\}", "Fig. 1", s)
    s = re.sub(r"Fig\.~\\ref\{fig:dvu\}", "Fig. 2", s)
    s = s.replace("5{,}", "5,").replace("{,}", ",").replace("\\,", " ").replace("~", " ").replace("\\ ", " ")
    s = s.replace("``", "“").replace("''", "”").replace("---", "—").replace("--", "–").replace("\\%", "%")
    s = s.replace("\\&", "&").replace("d'Armagnac", "d’Armagnac")
    s = re.sub(r"\\(textit|emph)\{([^{}]*)\}", "\x03\\2\x04", s)
    s = re.sub(r"\\textbf\{([^{}]*)\}", "\x05\\1\x06", s)
    s = re.sub(r"\\url\{([^}]*)\}", r"\1", s)
    s = re.sub(r"\\[a-zA-Z]+\*?", "", s).replace("{", "").replace("}", "")
    return re.sub(r"\s+", " ", s).strip()


def add_runs(p, s, size=10, base_italic=False, base_bold=False):
    it, bd, sub = base_italic, base_bold, False
    buf = ""

    def flush():
        nonlocal buf
        if buf:
            r = p.add_run(buf)
            r.italic, r.bold, r.font.size = it, bd, Pt(size)
            r.font.subscript = sub
            buf = ""
    for ch in s:
        if ch in "\x01\x02\x03\x04\x05\x06":
            flush()
            sub = True if ch == "\x01" else False if ch == "\x02" else sub
            it = True if ch == "\x03" else base_italic if ch == "\x04" else it
            bd = True if ch == "\x05" else base_bold if ch == "\x06" else bd
        else:
            buf += ch
    flush()


doc = Document()
st = doc.styles["Normal"]
st.font.name, st.font.size = "Times New Roman", Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(0)
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.625)
sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(1.0)


def para(text="", size=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY, italic=False, bold=False, indent=True, after=0, before=0):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.first_line_indent = Inches(0.14) if indent else None
    pf.space_after, pf.space_before = Pt(after), Pt(before)
    add_runs(p, text, size, italic, bold)
    return p


title = clean(re.search(r"\\title\{(.+?)\}\n", tex).group(1))
para(title, 24, WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=12)
for name, aff in re.findall(r"\\IEEEauthorblockN\{([^}]+)\}\s*\\IEEEauthorblockA\{(.+?)\}\}?\s*(?:\\and|\n\\maketitle|\n)", tex, re.S):
    para(name, 11, WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    for line in aff.split("\\\\"):
        para(clean(line), 10, WD_ALIGN_PARAGRAPH.CENTER, italic=True, indent=False)
    para("", 6, indent=False)

# two-column body
body = doc.add_section(WD_SECTION.CONTINUOUS)
cols = body._sectPr.xpath("./w:cols")
c = cols[0] if cols else OxmlElement("w:cols")
c.set(qn("w:num"), "2")
c.set(qn("w:space"), "360")
if not cols:
    body._sectPr.append(c)

abstract = clean(re.search(r"\\begin\{abstract\}(.+?)\\end\{abstract\}", tex, re.S).group(1))
p = para("", 9, indent=True)
add_runs(p, "\x03Abstract\x04—", 9, base_bold=True)
add_runs(p, abstract, 9, base_bold=True)
kw = clean(re.search(r"\\begin\{IEEEkeywords\}(.+?)\\end\{IEEEkeywords\}", tex, re.S).group(1))
p = para("", 9, after=6)
add_runs(p, "\x03Index Terms\x04—", 9, base_bold=True)
add_runs(p, kw, 9, base_bold=True)

main = tex.split("\\end{IEEEkeywords}", 1)[1].split("\\bibliographystyle", 1)[0]
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
sec_n, sub_n = 0, 0
fig_n = 0
blocks = re.split(r"\n\s*\n", main)
for b in blocks:
    b = b.strip()
    if not b:
        continue
    if b.startswith("\\begin{table}"):
        cap = clean(re.search(r"\\caption\{(.+?)\}\n", b).group(1))
        para("TABLE I", 8, WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=6)
        para(cap.upper(), 8, WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=3)
        rows = [r for r in re.search(r"\\midrule|\\toprule", b) and re.findall(r"^(.*?)\\\\\s*$", b, re.M)]
        rows = [[clean(c) for c in r.split("&")] for r in rows if "&" in r]
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = "Table Grid"
        t.autofit = False
        for j, wdt in enumerate([1.45, 0.65, 0.6, 0.6]):
            for cell in t.columns[j].cells:
                cell.width = Inches(wdt)
        for i, r in enumerate(rows):
            for j, val in enumerate(r):
                cell = t.cell(i, j)
                cp = cell.paragraphs[0]
                cp.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
                add_runs(cp, val, 8, base_bold=(i == 0))
        para("", 6, indent=False)
        continue
    if b.startswith("\\begin{figure}"):
        fig_n += 1
        img = re.search(r"figures/([^}]+)\.pdf", b).group(1)
        cap = clean(re.search(r"\\caption\{(.+?)\}\n", b, re.S).group(1))
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(FIG / f"{img}.png"), width=Inches(3.45))
        para(f"Fig. {fig_n}. {cap}", 8, WD_ALIGN_PARAGRAPH.JUSTIFY, indent=False, after=6)
        continue
    m = re.match(r"\\section\*?\{(.+?)\}\s*(.*)", b, re.S)
    if m:
        star = b.startswith("\\section*")
        if not star:
            sec_n += 1
            sub_n = 0
        head = clean(m.group(1))
        para(head.upper() if star else f"{ROMAN[sec_n - 1]}. {head.upper()}", 10, WD_ALIGN_PARAGRAPH.CENTER,
             indent=False, before=8, after=4)
        b = m.group(2).strip()
        if not b:
            continue
    m = re.match(r"\\subsection\{(.+?)\}\s*(.*)", b, re.S)
    if m:
        sub_n += 1
        para(f"{chr(64 + sub_n)}. {clean(m.group(1))}", 10, WD_ALIGN_PARAGRAPH.LEFT, italic=True, indent=False,
             before=6, after=3)
        b = m.group(2).strip()
        if not b:
            continue
    if b.startswith("\\begin{enumerate}") or "\\begin{enumerate}" in b:
        pre, rest = b.split("\\begin{enumerate}", 1)
        if pre.strip():
            para(clean(pre))
        items, post = rest.split("\\end{enumerate}", 1)
        for i, it in enumerate([x for x in items.split("\\item") if x.strip()], 1):
            p = para(f"{i}) {clean(it)}", indent=False)
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.first_line_indent = Inches(-0.15)
        if post.strip():
            para(clean(post))
        continue
    if b.startswith("\\begin{equation}"):
        para("R\x01eff\x02 = κ\x01R\x02 R / max(w, w\x01min\x02),    w\x01min\x02 = 0.02          (1)", 10,
             WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=3, after=3)
        continue
    if "\\begin{equation}" in b:
        pre, rest = b.split("\\begin{equation}", 1)
        para(clean(pre))
        para("R\x01eff\x02 = κ\x01R\x02 R / max(w, w\x01min\x02),    w\x01min\x02 = 0.02          (1)", 10,
             WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=3, after=3)
        post = rest.split("\\end{equation}", 1)[1]
        if post.strip():
            para(clean(post), indent=False)
        continue
    if b.startswith("\\end{document}") or b.startswith("\\bibliography"):
        continue
    para(clean(b))

# references from the IEEEtran .bbl
para("REFERENCES", 10, WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=8, after=4)
bbl = (HERE / "main.bbl").read_text(encoding="utf-8")
for key, entry in re.findall(r"\\bibitem\{([^}]+)\}\s*(.+?)(?=\\bibitem|\\end\{thebibliography\})", bbl, re.S):
    e = re.sub(r"\\BIBentry\w+", "", entry)
    e = e.replace("\\newblock", "")
    p = para(f"[{num[key]}] {clean(e)}", 8, indent=False)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)

out = HERE / "main.docx"
doc.save(out)
print(out)
